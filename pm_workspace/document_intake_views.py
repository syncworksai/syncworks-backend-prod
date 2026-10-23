"""Conservative document intake for existing property documents.

PDF (selectable-text), DOCX and TXT only. No OCR; suggestions are NOT trusted
data and cannot change a lease, rent, tenant, deposit or third-party account.
"""
from __future__ import annotations
import io
import re
from decimal import Decimal, InvalidOperation

from django.db import transaction
from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response

from .billing_views import _workspace
from .document_models import PMPropertyDocument
from .models import PMProperty

MAX_BYTES=8*1024*1024
FIELD_NAMES={"monthly_rent", "security_deposit", "lease_start", "lease_end"}
MONEY=r"\$?\s*([0-9]{1,3}(?:,[0-9]{3})+|[0-9]{1,7})(?:\.([0-9]{2}))?"
PATTERNS={
    "monthly_rent": r"(?i)\b(?:monthly\s+rent|total\s+monthly\s+rent|contract\s+rent|rent\s+to\s+owner|monthly\s+payment)\s*[:\-]?\s*"+MONEY,
    "security_deposit": r"(?i)\b(?:security\s+deposit|refundable\s+deposit)\s*[:\-]?\s*"+MONEY,
}
DATE_TEXT=r"(\d{1,2}/\d{1,2}/\d{4}|\d{4}-\d{2}-\d{2})"
DATE_PATTERNS={
    "lease_start":r"(?i)\b(?:lease\s+start|commencement\s+date|beginning\s+of\s+(?:the\s+)?lease|initial\s+lease\s+term\s+starts?)\s*[:\-]?\s*"+DATE_TEXT,
    "lease_end":r"(?i)\b(?:lease\s+end|lease\s+expiration|termination\s+date|initial\s+lease\s+term\s+ends?)\s*[:\-]?\s*"+DATE_TEXT,
}
def _parse_text(data,ext):
    if ext==".pdf":
        from pypdf import PdfReader
        reader=PdfReader(io.BytesIO(data),strict=False)
        if len(reader.pages)>30:
            return None,"PDF is longer than 30 pages. Split it into smaller documents first."
        pages=[(p.extract_text(extraction_mode="plain") or "") for p in reader.pages]
        return pages[:30],None
    if ext==".docx":
        from docx import Document
        doc=Document(io.BytesIO(data))
        texts=[p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                texts.extend(cell.text for cell in row.cells)
        return ["\n".join(texts)[:180000]],None
    if ext==".txt":
        return [data.decode("utf-8",errors="replace")[:180000]],None
    return None,"Upload a selectable-text PDF, DOCX or TXT file."

def _suggestions(pages):
    suggestions={}
    for field,pattern in {**PATTERNS,**DATE_PATTERNS}.items():
        matches=[]
        for i,page in enumerate(pages):
            for found in list(re.finditer(pattern,page))[:12]:
                raw=found.group(1).strip()
                if field in PATTERNS:
                    try:
                        amount=Decimal(raw.replace(",","") + "." + (found.group(2) or "00"))
                        if amount<=0 or amount>999999:continue
                        value=f"{amount:.2f}"
                    except (InvalidOperation,ValueError):
                        continue
                else:
                    from datetime import date,datetime
                    try:
                        dt=datetime.strptime(raw,"%m/%d/%Y").date() if "/" in raw else date.fromisoformat(raw)
                        value=dt.isoformat()
                    except ValueError: continue
                matches.append({"value":value,"page":i+1,
                                "context":re.sub(r"\s+"," ",page[max(0,found.start()-35):min(len(page),found.end()+35)]).strip()[:130]})
        if matches:
            values={m["value"] for m in matches}
            suggestions[field]={"value":matches[0]["value"],"candidates":matches[:4],
                                "confidence":"MULTIPLE_CANDIDATES" if len(values)>1 else "REVIEW_REQUIRED"}
    return suggestions

@api_view(["GET","POST","PATCH"])
@permission_classes([IsAuthenticated])
@transaction.atomic
def property_document_intake(request,property_id,document_id):
    workspace=_workspace(request)
    if not PMProperty.objects.filter(workspace=workspace,pk=property_id).exists():
        return Response({"detail":"Property not found in this portfolio."},status=404)
    doc=PMPropertyDocument.objects.select_for_update().filter(workspace=workspace,property_id=property_id,pk=document_id).first()
    if not doc:
        return Response({"detail":"Document not found."},status=404)
    state=dict(doc.extracted_terms or {})
    if request.method=="GET":
        return Response({"document_id":doc.id,"suggestions":state.get("intake_suggestions") or {},
                         "verified_fields":state.get("intake_verified") or {},
                         "reviewed_at":state.get("intake_reviewed_at")})
    if request.method=="PATCH":
        incoming=request.data.get("verified_fields")
        if not isinstance(incoming,dict) or set(incoming)-FIELD_NAMES:
            return Response({"detail":"Select only supported extracted fields."},status=400)
        recorded={}
        for key,value in incoming.items():
            if value in ("",None):continue
            if key in PATTERNS:
                try:
                    x=Decimal(str(value)).quantize(Decimal("0.01"))
                    if not x.is_finite() or x<0 or x>999999:
                        raise ValueError()
                    recorded[key]=str(x)
                except (InvalidOperation,ValueError,TypeError):
                    return Response({"detail":f"Invalid {key} amount."},status=400)
            else:
                from datetime import date
                try:
                    x=date.fromisoformat(str(value))
                    recorded[key]=x.isoformat()
                except ValueError:return Response({"detail":f"Invalid {key} date."},status=400)
        from django.utils import timezone
        state["intake_verified"]=recorded
        state["intake_reviewed_at"]=timezone.now().isoformat()
        state["intake_reviewed_by"]=request.user.pk
        doc.extracted_terms=state
        doc.save(update_fields=["extracted_terms","updated_at"])
        return Response({"detail":"Verified fields saved with document. Lease, tenant and ledger remain unchanged.",
                         "verified_fields":recorded})
    if not doc.document:
        return Response({"detail":"Upload the original document file first; a URL alone is not extracted."},status=400)
    filename=doc.document.name.lower()
    ext=next((x for x in (".pdf",".docx",".txt") if filename.endswith(x)),None)
    if not ext:
        return Response({"detail":"This document format is not supported for text extraction."},status=400)
    try:
        if doc.document.size>MAX_BYTES:
            return Response({"detail":"File exceeds the 8 MB extraction limit."},status=400)
        with doc.document.open("rb") as f:
            content=f.read(MAX_BYTES+1)
        if len(content)>MAX_BYTES:
            return Response({"detail":"File is too large to extract."},status=400)
        pages,reason=_parse_text(content,ext)
        if reason:return Response({"detail":reason},status=400)
        if not any(x.strip() for x in pages):
            return Response({"detail":"No selectable text was found. This looks scanned; manually enter values or provide a text PDF."},status=422)
        suggestions=_suggestions(pages)
    except Exception:
        return Response({"detail":"Could not read this file. Verify it is a valid, unencrypted document."},status=422)
    state["intake_suggestions"]=suggestions
    state.pop("intake_reviewed_at",None)
    doc.extracted_terms=state
    doc.save(update_fields=["extracted_terms","updated_at"])
    return Response({"document_id":doc.id,"suggestions":suggestions,"pages_reviewed":len(pages),
                     "notice":"Extraction is a draft. Review against the original before using any field; no profile, rent or ledger was modified."})
